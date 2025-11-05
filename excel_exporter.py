"""
Excel Exporter - Creates Excel files with invoice booking data
"""
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from config import Config


class ExcelExporter:
    """Handles Excel export for invoice bookings"""

    def __init__(self):
        """Initialize Excel exporter"""
        self.headers = Config.EXCEL_HEADERS

    def export_invoices(
        self,
        invoices: List[Dict[str, Optional[str]]],
        filename: Optional[str] = None
    ) -> Path:
        """
        Export invoice data to Excel file

        Args:
            invoices: List of invoice data dictionaries
            filename: Optional custom filename

        Returns:
            Path to the created Excel file
        """
        # Generate filename if not provided
        if not filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"KERN1_Invoices_{timestamp}.xlsx"

        # Ensure .xlsx extension
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'

        # Create full path
        filepath = Config.EXPORT_DIR / filename

        # Create DataFrame
        df = self._create_dataframe(invoices)

        # Export to Excel
        df.to_excel(filepath, index=False, sheet_name='Invoices')

        # Apply formatting
        self._apply_formatting(filepath)

        return filepath

    def _create_dataframe(
        self,
        invoices: List[Dict[str, Optional[str]]]
    ) -> pd.DataFrame:
        """
        Create pandas DataFrame from invoice data

        Args:
            invoices: List of invoice dictionaries

        Returns:
            Formatted DataFrame
        """
        # Convert to DataFrame
        df = pd.DataFrame(invoices)

        # Rename columns using headers mapping
        df = df.rename(columns=self.headers)

        # Ensure columns are in correct order
        ordered_columns = [self.headers[field] for field in Config.INVOICE_FIELDS]
        df = df[ordered_columns]

        # Format date column
        date_col = self.headers['invoice_date']
        if date_col in df.columns:
            df[date_col] = pd.to_datetime(df[date_col], errors='coerce')

        # Format amount columns
        for field in ['total_amount', 'vat_amount']:
            col = self.headers[field]
            if col in df.columns:
                df[col] = pd.to_numeric(df[col], errors='coerce')

        return df

    def _apply_formatting(self, filepath: Path):
        """
        Apply professional formatting to Excel file

        Args:
            filepath: Path to Excel file
        """
        try:
            # Load workbook
            wb = load_workbook(filepath)
            ws = wb.active

            # Define styles
            header_fill = PatternFill(
                start_color="366092",
                end_color="366092",
                fill_type="solid"
            )
            header_font = Font(
                bold=True,
                color="FFFFFF",
                size=11
            )
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # Format headers
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border

            # Format data rows
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical='center')

            # Adjust column widths
            column_widths = {
                'A': 15,  # Invoice Date
                'B': 18,  # Invoice Number
                'C': 40,  # Service Description
                'D': 25,  # Supplier Name
                'E': 15,  # Total Amount
                'F': 15,  # VAT Amount
            }

            for col, width in column_widths.items():
                ws.column_dimensions[col].width = width

            # Format date column (A)
            for cell in ws['A'][1:]:
                if cell.value:
                    cell.number_format = 'YYYY-MM-DD'

            # Format amount columns (E, F)
            for col in ['E', 'F']:
                for cell in ws[col][1:]:
                    if cell.value:
                        cell.number_format = '€#,##0.00'

            # Freeze header row
            ws.freeze_panes = 'A2'

            # Save formatted workbook
            wb.save(filepath)

        except Exception as e:
            print(f"Warning: Could not apply formatting: {e}")

    def append_to_existing(
        self,
        invoices: List[Dict[str, Optional[str]]],
        existing_file: Path
    ) -> Path:
        """
        Append new invoices to existing Excel file

        Args:
            invoices: List of new invoice data
            existing_file: Path to existing Excel file

        Returns:
            Path to the updated Excel file
        """
        # Read existing data
        existing_df = pd.read_excel(existing_file)

        # Create new data DataFrame
        new_df = self._create_dataframe(invoices)

        # Combine
        combined_df = pd.concat([existing_df, new_df], ignore_index=True)

        # Remove duplicates based on invoice number
        invoice_col = self.headers['invoice_number']
        combined_df = combined_df.drop_duplicates(subset=[invoice_col], keep='last')

        # Sort by date
        date_col = self.headers['invoice_date']
        combined_df = combined_df.sort_values(by=date_col, ascending=False)

        # Export
        combined_df.to_excel(existing_file, index=False, sheet_name='Invoices')

        # Apply formatting
        self._apply_formatting(existing_file)

        return existing_file

    def create_summary(self, invoices: List[Dict[str, Optional[str]]]) -> Dict:
        """
        Create summary statistics for invoices

        Args:
            invoices: List of invoice data

        Returns:
            Dictionary with summary statistics
        """
        df = self._create_dataframe(invoices)

        total_col = self.headers['total_amount']
        vat_col = self.headers['vat_amount']

        summary = {
            'total_invoices': len(df),
            'total_amount': df[total_col].sum() if total_col in df else 0,
            'total_vat': df[vat_col].sum() if vat_col in df else 0,
            'earliest_date': df[self.headers['invoice_date']].min() if len(df) > 0 else None,
            'latest_date': df[self.headers['invoice_date']].max() if len(df) > 0 else None,
        }

        return summary
